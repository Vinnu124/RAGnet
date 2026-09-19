"""Document → per-page markdown. Keeps headings, tables and page numbers (the old code threw all of that away)."""
from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".pdf", ".md", ".markdown", ".txt", ".html", ".htm", ".docx"}


@dataclass
class Page:
    number: int  # 1-based
    text: str  # markdown


@dataclass
class ParsedDoc:
    name: str
    path: Path
    sha256: str
    pages: list[Page] = field(default_factory=list)

    @property
    def doc_id(self) -> str:
        return self.sha256[:12]


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


# ----------------------------------------------------------------------------- PDF
def _parse_pdf_fast(path: Path) -> list[Page]:
    """pymupdf4llm: very fast, markdown with headings/tables, page-aware."""
    import pymupdf4llm

    chunks = pymupdf4llm.to_markdown(str(path), page_chunks=True, show_progress=False)
    pages = []
    for c in chunks:
        text = (c.get("text") or "").strip()
        if text:
            pages.append(Page(number=int(c["metadata"]["page"]), text=text))
    return pages


def _parse_pdf_hifi(path: Path) -> list[Page]:
    """Docling: layout model + table structure + OCR for scanned pages. Much slower, much better on hard PDFs."""
    try:
        from docling.document_converter import DocumentConverter
        from docling_core.types.doc import DocItemLabel, TableItem
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Install the hifi extra:  uv sync --extra hifi") from e

    result = DocumentConverter().convert(str(path))
    doc = result.document
    by_page: dict[int, list[str]] = {}
    for item, _level in doc.iterate_items():
        if not getattr(item, "prov", None):
            continue
        page_no = item.prov[0].page_no
        if isinstance(item, TableItem):
            md = item.export_to_markdown(doc)
        else:
            text = getattr(item, "text", "") or ""
            if not text.strip():
                continue
            label = getattr(item, "label", None)
            if label == DocItemLabel.SECTION_HEADER:
                level = getattr(item, "level", 1) or 1
                md = f"{'#' * min(level + 1, 6)} {text}"
            elif label == DocItemLabel.TITLE:
                md = f"# {text}"
            elif label == DocItemLabel.LIST_ITEM:
                md = f"- {text}"
            else:
                md = text
        by_page.setdefault(page_no, []).append(md)
    return [Page(number=p, text="\n\n".join(parts)) for p, parts in sorted(by_page.items())]


# ----------------------------------------------------------------------------- others
def _parse_text(path: Path) -> list[Page]:
    return [Page(number=1, text=path.read_text(encoding="utf-8", errors="ignore").strip())]


def _parse_html(path: Path) -> list[Page]:
    import re

    html = path.read_text(encoding="utf-8", errors="ignore")
    html = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.DOTALL | re.IGNORECASE)
    for lvl in range(1, 7):
        html = re.sub(rf"<h{lvl}[^>]*>(.*?)</h{lvl}>", lambda m, l=lvl: f"\n\n{'#' * l} {m.group(1)}\n\n", html, flags=re.DOTALL | re.IGNORECASE)
    html = re.sub(r"</(p|div|li|tr|br)\s*>", "\n", html, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", html)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return [Page(number=1, text=text.strip())]


def _parse_docx(path: Path) -> list[Page]:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("pip install python-docx to ingest .docx files") from e
    d = docx.Document(str(path))
    parts = []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            lvl = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append(f"{'#' * min(int(lvl), 6)} {p.text}")
        else:
            parts.append(p.text)
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        if rows:
            parts.append("| " + rows[0] + " |\n|" + "---|" * len(t.rows[0].cells) + "\n" + "\n".join(f"| {r} |" for r in rows[1:]))
    return [Page(number=1, text="\n\n".join(parts))]


# ----------------------------------------------------------------------------- entry
def parse(path: Path, hifi: bool = False) -> ParsedDoc:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type {ext}: {path}")
    if ext == ".pdf":
        pages = _parse_pdf_hifi(path) if hifi else _parse_pdf_fast(path)
    elif ext in {".md", ".markdown", ".txt"}:
        pages = _parse_text(path)
    elif ext in {".html", ".htm"}:
        pages = _parse_html(path)
    else:
        pages = _parse_docx(path)
    return ParsedDoc(name=path.name, path=path, sha256=file_sha256(path), pages=pages)
